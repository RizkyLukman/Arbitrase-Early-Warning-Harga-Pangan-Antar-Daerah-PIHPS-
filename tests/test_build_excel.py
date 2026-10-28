"""Tes src/build_excel.py: struktur workbook, sel bernama, rumus yang didukung, bahasa dashboard."""
from __future__ import annotations

import re
from pathlib import Path

import pytest
from openpyxl import load_workbook

from src.audit_dashboard import ISTILAH_TERLARANG, PROCESSED, SHEET_DASHBOARD, nilai_acuan, pemetaan_acuan
from src.analysis.common import load_cfg
from src.build_excel import build

RUMUS_TERLARANG = ["FILTER(", "XLOOKUP(", "SORT(", "UNIQUE(", "SEQUENCE(", "XMATCH(", "LET("]   # butuh Excel 365 / tidak ada di LibreOffice


@pytest.fixture(scope="module")
def xlsx(tmp_path_factory) -> Path:
    if not (PROCESSED / "fact_harga.parquet").exists():
        pytest.skip("data/processed/ belum dibuat")
    return build(tmp_path_factory.mktemp("xl") / "dash.xlsx")


def test_sheets_in_order_and_data_sheets_present(xlsx: Path):
    wb = load_workbook(xlsx)
    assert wb.sheetnames[:7] == SHEET_DASHBOARD + ["Catatan metode"]
    assert {"d_komoditas", "d_kota", "d_rute", "d_spread", "d_gejolak", "d_uji_peringatan", "d_tindakan", "d_hal", "d_harga_tujuan"} <= set(wb.sheetnames)


def test_every_audited_name_exists(xlsx: Path):
    wb = load_workbook(xlsx)
    cfg, lcfg = load_cfg()
    peta = pemetaan_acuan(nilai_acuan(cfg, lcfg))
    hilang = [n for n in peta if n not in wb.defined_names]
    assert hilang == []


def test_formulas_use_only_widely_supported_functions(xlsx: Path):
    wb = load_workbook(xlsx)
    salah = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("=") and any(f in c.value.upper() for f in RUMUS_TERLARANG):
                    salah.append(f"{ws.title}!{c.coordinate}")
    assert salah == []


def test_dashboard_text_has_no_technical_terms(xlsx: Path):
    wb = load_workbook(xlsx)
    temuan = []
    for sheet in SHEET_DASHBOARD:
        for row in wb[sheet].iter_rows(max_col=14):
            for c in row:
                if isinstance(c.value, str) and not c.value.startswith("=") and any(k in c.value.lower() for k in ISTILAH_TERLARANG):
                    temuan.append(f"{sheet}!{c.coordinate}: {c.value[:60]}")
    assert temuan == []


def test_inputs_have_validation_and_decisions_exist(xlsx: Path):
    wb = load_workbook(xlsx)
    assert any("B5" in str(dv.sqref) for dv in wb["Gejolak harga"].data_validations.dataValidation)
    assert any("B5" in str(dv.sqref) for dv in wb["Pindah stok"].data_validations.dataValidation)
    for sheet in SHEET_DASHBOARD:
        teks = [c.value for row in wb[sheet].iter_rows(max_col=1) for c in row if isinstance(c.value, str)]
        assert any(t.startswith("KEPUTUSAN") for t in teks), sheet
        assert sum(bool(re.match(r"^[123]\. ", t)) for t in teks) >= 3, sheet
