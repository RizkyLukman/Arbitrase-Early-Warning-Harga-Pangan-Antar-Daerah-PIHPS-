"""Tes src/build_excel_desain.py: salinan dashboard bergaya mockup tetap bisa diaudit dan bahasanya bersih."""
from __future__ import annotations

import re
import zipfile
from pathlib import Path

import pytest
from openpyxl import load_workbook

from src.analysis.common import load_cfg
from src.audit_dashboard import ISTILAH_TERLARANG, PROCESSED, SHEET_DASHBOARD, nilai_acuan, pemetaan_acuan
from src.build_excel_desain import build

RUMUS_TERLARANG = ["FILTER(", "XLOOKUP(", "SORT(", "UNIQUE(", "SEQUENCE(", "XMATCH(", "LET("]
SIMBOL_TERLARANG = ["→", "·", "—", "×"]   # permintaan pemilik: bahasa analis, bukan simbol


@pytest.fixture(scope="module")
def xlsx(tmp_path_factory) -> Path:
    if not (PROCESSED / "fact_harga.parquet").exists():
        pytest.skip("data/processed/ belum dibuat")
    return build(tmp_path_factory.mktemp("xl") / "desain.xlsx")


def _teks_tampil(wb):
    """Semua teks di sheet dashboard (nilai tetap + teks literal di dalam rumus), kolom isi saja (bukan kolom bantu)."""
    for name in SHEET_DASHBOARD + ["Catatan metode"]:
        for row in wb[name].iter_rows(max_col=27):
            for c in row:
                v = c.value
                if v is None:
                    continue
                v = str(v)
                yield name, c.coordinate, " ".join(re.findall(r'"([^"]*)"', v)) if v.startswith("=") else v


def test_sheets_in_menu_order(xlsx: Path):
    wb = load_workbook(xlsx)
    assert wb.sheetnames[:7] == SHEET_DASHBOARD + ["Catatan metode"]


def test_every_audited_name_exists(xlsx: Path):
    wb = load_workbook(xlsx)
    cfg, lcfg = load_cfg()
    hilang = [n for n in pemetaan_acuan(nilai_acuan(cfg, lcfg)) if n not in wb.defined_names]
    assert hilang == []


def test_formulas_use_only_widely_supported_functions(xlsx: Path):
    wb = load_workbook(xlsx)
    salah = [f"{ws.title}!{c.coordinate}" for ws in wb.worksheets for row in ws.iter_rows() for c in row
             if isinstance(c.value, str) and c.value.startswith("=") and any(f in c.value.upper() for f in RUMUS_TERLARANG)]
    assert salah == []


def test_no_technical_terms_or_ai_symbols_on_screen(xlsx: Path):
    wb = load_workbook(xlsx, rich_text=True)
    istilah, simbol = [], []
    for name, ref, teks in _teks_tampil(wb):
        t = teks.lower()
        if name != "Catatan metode" and any(k in t for k in ISTILAH_TERLARANG):
            istilah.append(f"{name}!{ref}: {teks[:60]}")
        if any(s in teks for s in SIMBOL_TERLARANG):
            simbol.append(f"{name}!{ref}: {teks[:60]}")
    assert istilah == []
    assert simbol == []


def test_map_picture_is_drawn_below_bubble_chart(xlsx: Path):
    with zipfile.ZipFile(xlsx) as z:
        for n in z.namelist():
            if n.startswith("xl/drawings/drawing") and n.endswith(".xml"):
                xml = z.read(n).decode()
                if ":pic>" in xml or "<pic>" in xml:
                    pic = min(xml.find(t) for t in ("<xdr:pic>", "<pic>") if xml.find(t) >= 0)
                    frame = min(xml.find(t) for t in ("<xdr:graphicFrame", "<graphicFrame") if xml.find(t) >= 0)
                    assert pic < frame
                    return
    pytest.fail("gambar peta tidak ditemukan")


def test_every_sheet_ends_with_decision_band(xlsx: Path):
    wb = load_workbook(xlsx)
    for name in SHEET_DASHBOARD:
        assert any(c.value == "FAKTA DAN KEPUTUSAN" for row in wb[name].iter_rows() for c in row), name
