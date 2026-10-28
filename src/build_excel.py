"""Tahap 6 — Membangun dashboard Excel dari data/processed/ (tanpa Power BI).

Output: excel/dashboard_harga_pangan.xlsx — 7 sheet dashboard + sheet data (d_*). Interaktif lewat sel input (dropdown
komoditas, pengali ongkos kirim, tambahan ongkos) dan rumus Excel klasik (SUMIFS, INDEX/MATCH, LARGE, SUMPRODUCT) sehingga
terbaca di Excel 2016+ dan bisa dihitung ulang. Angka hasil analisis (median per rute, gejolak per kota, uji peringatan, dll.)
ditulis sebagai DATA di sheet d_* (sumber: tabel data/processed, lihat sheet "Catatan metode"); rumus di sheet dashboard hanya
memilih/menyaring/mengurutkan data itu. Setiap angka kunci diberi NAMA (Formulas > Name Manager) agar bisa diaudit oleh
src/audit_dashboard.py.

    python -m src.build_excel            # tulis excel/dashboard_harga_pangan.xlsx
    python -m src.build_excel --out X    # ke path lain (dipakai tes)
"""
from __future__ import annotations

import argparse
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, BubbleChart, LineChart, Reference, Series
from openpyxl.chart.axis import DateAxis
from openpyxl.drawing.text import CharacterProperties, Font as DFont, ParagraphProperties
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from src.analysis.common import KOMODITAS_URUT, LABEL_PENDEK, ROOT, load_cfg

PROCESSED = ROOT / "data" / "processed"
OUT = ROOT / "excel" / "dashboard_harga_pangan.xlsx"
FONT = "Arial"
TEAL, BLUE, ORANGE, ORANGE_TXT, GRAY, INK, MUTED, DARK, BG, WHITE = "0B7A57", "2A78D6", "EB6834", "C2410C", "B4BDB9", "15201D", "6B7A75", "0F2A26", "F3F5F4", "FFFFFF"
RP = '"Rp "#,##0'
PCT = "0%"
BULAN_SUMBU = "[$-421]mmm yy"   # 421 = Indonesia: nama bulan Indonesia di sumbu apa pun bahasa Excel-nya
BULAN = ["Jan", "Feb", "Mar", "Apr", "Mei", "Jun", "Jul", "Agu", "Sep", "Okt", "Nov", "Des"]
NAMA_SHEET = {"ringkasan": "Ringkasan", "gejolak": "Gejolak harga", "rute": "Pindah stok", "peringatan": "Peringatan dini",
              "pola": "Pola musiman", "rantai": "Rantai pasok", "metode": "Catatan metode"}


def q(sheet: str) -> str:
    return f"'{sheet}'" if " " in sheet else sheet


# ---------------------------------------------------------------- gaya
def font(size=10, bold=False, color=INK, italic=False) -> Font:
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


def fill(color: str) -> PatternFill:
    return PatternFill("solid", start_color=color, end_color=color)


def set_widths(ws, widths: dict[str, float]) -> None:
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def teks(ws, ref: str, value, size=10, bold=False, color=INK, wrap=False, align="left", fmt=None, bg=None, italic=False):
    c = ws[ref]
    c.value = value
    c.font = font(size, bold, color, italic)
    c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    if fmt:
        c.number_format = fmt
    if bg:
        c.fill = fill(bg)
    return c


def kartu(ws, col: int, row: int, label: str, value, fmt: str | None = None, color: str = INK, lebar: int = 3, sub: str | None = None):
    """Kartu KPI: label kecil, angka besar, keterangan kecil; menempati `lebar` kolom x 3 baris."""
    c1, c2 = get_column_letter(col), get_column_letter(col + lebar - 1)
    for r in range(row, row + 3):
        for cc in range(col, col + lebar):
            ws.cell(r, cc).fill = fill(WHITE)
    ws.merge_cells(f"{c1}{row}:{c2}{row}")
    ws.merge_cells(f"{c1}{row + 1}:{c2}{row + 1}")
    ws.merge_cells(f"{c1}{row + 2}:{c2}{row + 2}")
    teks(ws, f"{c1}{row}", label.upper(), 8, True, MUTED, bg=WHITE)
    teks(ws, f"{c1}{row + 1}", value, 20, True, color, fmt=fmt, bg=WHITE)
    teks(ws, f"{c1}{row + 2}", sub or "", 8, False, MUTED, wrap=True, bg=WHITE)
    ws.row_dimensions[row + 1].height = 32
    ws.row_dimensions[row + 2].height = 26
    return f"{c1}{row + 1}"


def keputusan(ws, row: int, baris: list[tuple[str, str]], lebar_kolom: str = "N") -> None:
    """Band gelap di akhir sheet: 3 keputusan, masing-masing dengan fakta berangka yang tampil di sheet yang sama."""
    teks(ws, f"A{row}", "KEPUTUSAN (dan fakta yang mendasarinya)", 11, True, WHITE, bg=DARK)
    ws.merge_cells(f"A{row}:{lebar_kolom}{row}")
    for i, (kep, fakta) in enumerate(baris, 1):
        r = row + i
        ws.merge_cells(f"A{r}:{lebar_kolom}{r}")
        c = ws[f"A{r}"]
        c.value = f"{i}. {kep}   —   Fakta: {fakta}"
        c.font = font(10, True, WHITE)
        c.fill = fill(DARK)
        c.alignment = Alignment(vertical="center", wrap_text=True)
        for cc in range(1, 15):
            ws.cell(r, cc).fill = fill(DARK)
        ws.row_dimensions[r].height = 34


def nama(wb: Workbook, name: str, ref: str) -> None:
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


def tabel_data(wb: Workbook, sheet: str, df: pd.DataFrame, nama_tabel: str, formats: dict[str, str] | None = None) -> dict[str, str]:
    """Tulis DataFrame ke sheet data sebagai Excel Table. Mengembalikan {kolom: rentang absolut tanpa header}."""
    ws = wb.create_sheet(sheet)
    ws.sheet_properties.tabColor = GRAY
    ws.sheet_state = "hidden"   # data pendukung; tampilkan lewat klik kanan tab > Unhide
    for j, col in enumerate(df.columns, 1):
        teks(ws, f"{get_column_letter(j)}1", col, 9, True, WHITE, bg=DARK)
    for i, row in enumerate(df.itertuples(index=False), 2):
        for j, v in enumerate(row, 1):
            if pd.isna(v) if not isinstance(v, str) else False:
                v = None
            elif hasattr(v, "item"):
                v = v.item()
            c = ws.cell(i, j, v)
            c.font = font(9)
            if formats and df.columns[j - 1] in formats:
                c.number_format = formats[df.columns[j - 1]]
    n = len(df)
    ref = f"A1:{get_column_letter(len(df.columns))}{max(n + 1, 2)}"
    t = Table(displayName=nama_tabel, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws.add_table(t)
    for j, col in enumerate(df.columns, 1):
        ws.column_dimensions[get_column_letter(j)].width = max(10, min(40, len(str(col)) + 2))
    return {col: f"{q(sheet)}!${get_column_letter(j)}$2:${get_column_letter(j)}${n + 1}" for j, col in enumerate(df.columns, 1)}


def halaman_dasar(wb: Workbook, key: str, judul_formula_or_text, sub: str | None = None):
    ws = wb.create_sheet(NAMA_SHEET[key])
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    ws.sheet_properties.tabColor = TEAL
    set_widths(ws, {get_column_letter(i): 12 for i in range(1, 15)})
    ws.column_dimensions["A"].width = 24
    ws.row_dimensions[1].height = 8
    ws.row_dimensions[2].height = 46
    ws.merge_cells("A2:L2")
    teks(ws, "A2", judul_formula_or_text, 15, True, INK, wrap=True)
    if sub:
        ws.merge_cells("A3:L3")
        teks(ws, "A3", sub, 9, False, MUTED)
    for r in range(1, 60):
        for c in range(1, 16):
            if ws.cell(r, c).fill.fgColor.rgb in (None, "00000000"):
                ws.cell(r, c).fill = fill(BG)
    return ws


def rapikan(chart, judul_pt: int = 10) -> None:
    """Sumbu ditampilkan (openpyxl 3.1 menyembunyikannya secara default) dan data di kolom tersembunyi tetap digambar.
    Judul & legenda diberi ruang sendiri (overlay=False; tanpa ini Excel menumpuknya di atas area grafik), judul Arial kecil."""
    for ax in (chart.x_axis, chart.y_axis):
        ax.delete = False
    chart.visible_cells_only = False
    chart.style = 2
    if chart.title is not None:
        chart.title.overlay = False
        cp = CharacterProperties(sz=judul_pt * 100, b=True, solidFill=INK, latin=DFont(typeface=FONT))
        for p in chart.title.tx.rich.p:
            p.pPr = ParagraphProperties(defRPr=cp)
            for r in p.r or []:
                r.rPr = cp
    if chart.legend is not None:
        chart.legend.overlay = False
        chart.legend.position = "b"


def rp(x: float) -> str:
    return f"{x:,.0f}".replace(",", ".")


def tanpa_awalan(kalimat: str) -> str:
    """'cabai rawit merah: untung Rp ...' -> 'untung Rp ...' (alasan_singkat diawali nama komoditas)."""
    return kalimat.split(": ", 1)[1] if ": " in kalimat else kalimat


def sumbu_bulan(chart, tiap: int = 1) -> None:
    """Sumbu tanggal berlabel bulan (Indonesia) tiap `tiap` bulan. Tanpa ini Excel mengubah sumbu kategori bertanggal
    menjadi sumbu tanggal berlabel harian -> label tiap minggu berdesakan."""
    chart.y_axis.crossAx = 500
    chart.x_axis = DateAxis(crossAx=100)
    chart.x_axis.number_format = BULAN_SUMBU
    chart.x_axis.majorTimeUnit = "months"
    chart.x_axis.majorUnit = tiap


def rupiah_teks(expr: str) -> str:
    """Rumus teks angka bulat dengan titik ribuan, tidak bergantung pengaturan bahasa Excel (FIXED mengikuti locale)."""
    return f'SUBSTITUTE(FIXED({expr},0),",",".")'


# ---------------------------------------------------------------- data
def muat() -> dict[str, pd.DataFrame]:
    t = {p.stem: pd.read_parquet(p) for p in PROCESSED.glob("*.parquet")}
    kom = t["dim_komoditas"].set_index("komoditas_key")
    kota = t["dim_kota"].set_index("kota_key")
    for nm in ("fact_gejolak", "fact_penentu_arah", "fact_margin", "fact_tindakan", "fact_hal", "fact_pola_idul_fitri", "fact_alert", "fact_harga_mingguan"):
        if "komoditas_key" in t[nm]:
            t[nm]["komoditas"] = t[nm].komoditas_key.map(kom.nama_dashboard)
        if "kota_key" in t[nm]:
            t[nm]["kota"] = t[nm].kota_key.map(kota.kota)
    return t


def build(out: Path = OUT) -> Path:
    cfg, lcfg = load_cfg()
    T = muat()
    kom, kota, rute = T["dim_komoditas"], T["dim_kota"], T["dim_rute"]
    info = T["dim_info"].set_index("kunci").nilai
    akhir = int(T["fact_harga_mingguan"].minggu_key.max())
    minggu_akhir = pd.to_datetime(str(akhir), format="%Y%m%d")
    label_minggu = f"Minggu {(minggu_akhir - pd.Timedelta(days=4)).day}–{minggu_akhir.day} {BULAN[minggu_akhir.month - 1]} {minggu_akhir.year}"
    kons = lcfg["sensitivitas"]["skenario_konservatif"]
    tujuan = list(cfg["wilayah"]["kota_tujuan"])
    urut_kom = [LABEL_PENDEK[k] for k in KOMODITAS_URUT]

    wb = Workbook()
    wb.remove(wb.active)
    # sheet dashboard dibuat dulu (urutan tab), diisi setelah sheet data ada
    S = {k: halaman_dasar(wb, k, "") for k in ("ringkasan", "gejolak", "rute", "peringatan", "pola", "rantai")}
    ws_met = wb.create_sheet(NAMA_SHEET["metode"])

    # ---------- sheet data (angka hasil analisis; sumber data/processed)
    d_kom = kom[["komoditas_key", "nama_dashboard", "satuan", "peringkat_gejolak", "rata_perubahan_per_minggu", "gejolak_mingguan_median",
                 "idulfitri_mulai_naik_minggu", "idulfitri_kenaikan_minggu_min1", "idulfitri_tahun_konsisten", "idulfitri_puncak_minggu",
                 "idulfitri_puncak_kenaikan", "hemat_beli_produsen_pct", "hemat_beli_pedagang_besar_pct", "modern_vs_tradisional_pct"]].copy()
    d_kom = d_kom.set_index("nama_dashboard").loc[urut_kom].reset_index()
    R_kom = tabel_data(wb, "d_komoditas", d_kom, "d_komoditas", {"rata_perubahan_per_minggu": PCT, "gejolak_mingguan_median": PCT,
                       "idulfitri_kenaikan_minggu_min1": PCT, "idulfitri_puncak_kenaikan": PCT, "hemat_beli_produsen_pct": PCT,
                       "hemat_beli_pedagang_besar_pct": PCT, "modern_vs_tradisional_pct": PCT})
    d_kota = kota[["kota_key", "kota", "provinsi", "lat", "lon", "peran", "komoditas_sumber", "penentu_arah_untuk"]]
    R_kota = tabel_data(wb, "d_kota", d_kota, "d_kota")
    g = T["fact_gejolak"][["kota", "komoditas", "gejolak_mingguan", "rata_perubahan_per_minggu", "peringkat_dalam_komoditas"]]
    R_gej = tabel_data(wb, "d_gejolak", g, "d_gejolak", {"gejolak_mingguan": "0.0%", "rata_perubahan_per_minggu": PCT})
    d_rute = rute[["rute_key", "nama_rute", "komoditas_key", "jarak_jalan_km", "median_selisih_harga_rp_per_kg", "ongkos_angkut_rp_per_kg", "susut_pct",
                   "median_ongkos_rp_per_kg", "median_untung_rp_per_kg", "porsi_minggu_untung", "minggu_untung_dari_10", "rata_bertahan_minggu",
                   "minggu_untung_berturut_terakhir", "layak", "layak_jika_jual_ke_pedagang_besar", "median_untung_jual_ke_pedagang_besar", "status"]].copy()
    d_rute["risiko"] = d_rute.apply(lambda r: "aman" if r.layak and r.layak_jika_jual_ke_pedagang_besar else ("hanya eceran" if r.layak else "tidak layak"), axis=1)
    R_rute = tabel_data(wb, "d_rute", d_rute, "d_rute", {"median_selisih_harga_rp_per_kg": RP, "median_ongkos_rp_per_kg": RP, "median_untung_rp_per_kg": RP,
                        "ongkos_angkut_rp_per_kg": RP, "porsi_minggu_untung": PCT, "median_untung_jual_ke_pedagang_besar": RP})
    sens = T["fact_rute_sensitivitas"].copy()
    R_sens = tabel_data(wb, "d_sensitivitas", sens, "d_sensitivitas", {"median_untung_rp_per_kg": RP, "porsi_minggu_untung": PCT})
    sp = T["fact_spread"][["rute_key", "minggu_key", "harga_sumber", "harga_tujuan", "selisih_harga", "ongkos_angkut", "susut_rp"]].copy()
    sp["minggu"] = pd.to_datetime(sp.minggu_key.astype(str), format="%Y%m%d")
    sp["kunci"] = sp.rute_key.astype(str) + "|" + sp.minggu_key.astype(str)
    R_sp = tabel_data(wb, "d_spread", sp, "d_spread", {"minggu": "yyyy-mm-dd"})
    pa = T["fact_penentu_arah"]
    pa = pa[pa.tampil_di_dashboard].sort_values(["komoditas_key", "peringkat_penentu_arah"])
    pa = pa.assign(urut=pa.komoditas.map({k: i for i, k in enumerate(urut_kom)})).sort_values(["urut", "peringkat_penentu_arah"])
    R_pa = tabel_data(wb, "d_penentu_arah", pa[["komoditas", "kota", "jeda_minggu", "korelasi", "hubungan"]], "d_penentu_arah", {"korelasi": "0.00"})
    uji = T["fact_uji_peringatan"].sort_values("urutan")
    R_uji = tabel_data(wb, "d_uji_peringatan", uji[["urutan", "aturan", "benar_dari_10", "tertangkap_dari_10", "lebih_awal_minggu", "porsi_minggu_peringatan", "jumlah_peringatan", "jumlah_lonjakan"]],
                       "d_uji_peringatan", {"porsi_minggu_peringatan": PCT})
    al = T["fact_alert"]
    al = al[(al.kota_key == 34) & (al.komoditas_key == 16) & al.harga.notna()].sort_values("minggu_key")
    al = pd.DataFrame({"minggu": pd.to_datetime(al.minggu_key.astype(str), format="%Y%m%d"), "harga": al.harga,
                       "harga_saat_peringatan": al.harga.where(al.peringatan), "awal_lonjakan": al.awal_lonjakan.fillna(False).astype(bool), "alasan": al.alasan})
    R_al = tabel_data(wb, "d_peringatan_jakarta", al, "d_peringatan_jakarta", {"minggu": "yyyy-mm-dd", "harga": "#,##0", "harga_saat_peringatan": "#,##0"})
    pi = T["fact_pola_idul_fitri"]
    pola_rows = []
    for k in urut_kom:
        tb = pi[pi.komoditas == k].pivot(index="minggu_ke_idul_fitri", columns="tahun_idul_fitri", values="kenaikan_vs_harga_normal").reindex(range(-8, 5))
        for m in tb.index:
            pola_rows.append({"komoditas": k, "minggu_ke_lebaran": int(m), "rata_rata": tb.loc[m].mean(), "terendah": tb.loc[m].min(), "tertinggi": tb.loc[m].max(),
                              **{f"th_{int(c)}": tb.loc[m, c] for c in tb.columns}})
    pola = pd.DataFrame(pola_rows)
    R_pola = tabel_data(wb, "d_pola_lebaran", pola, "d_pola_lebaran", {c: PCT for c in pola.columns if c not in ("komoditas", "minggu_ke_lebaran")})
    m = T["fact_margin"]
    m = m.assign(urut=m.komoditas.map({k: i for i, k in enumerate(urut_kom)})).sort_values(["urut", "hemat_beli_produsen_pct"], ascending=[True, False])
    R_m = tabel_data(wb, "d_margin", m[["komoditas", "kota", "cakupan", "harga_produsen_dari", "n_bulan", "harga_produsen", "harga_pedagang_besar", "harga_pasar",
                                        "tambahan_di_pedagang_besar_pct", "tambahan_di_pasar_pct", "hemat_beli_produsen_pct", "hemat_beli_pedagang_besar_pct"]],
                     "d_margin", {"harga_produsen": RP, "harga_pedagang_besar": RP, "harga_pasar": RP, "tambahan_di_pedagang_besar_pct": PCT,
                                  "tambahan_di_pasar_pct": PCT, "hemat_beli_produsen_pct": PCT, "hemat_beli_pedagang_besar_pct": PCT})
    marg_kom = (m.groupby("komoditas").agg(kota=("kota", "size"), harga_produsen=("harga_produsen", "median"), harga_pedagang_besar=("harga_pedagang_besar", "median"),
                                           harga_pasar=("harga_pasar", "median"), margin_pedagang_besar=("tambahan_di_pedagang_besar_pct", "median"),
                                           margin_pasar=("tambahan_di_pasar_pct", "median"), hemat_produsen=("hemat_beli_produsen_pct", "median"),
                                           hemat_pedagang_besar=("hemat_beli_pedagang_besar_pct", "median")).reindex([k for k in urut_kom if k != "Bawang putih"]).reset_index())
    marg_kom["tambahan_pedagang_besar_rp"] = marg_kom.harga_pedagang_besar - marg_kom.harga_produsen
    marg_kom["tambahan_pasar_rp"] = marg_kom.harga_pasar - marg_kom.harga_pedagang_besar
    marg_kom["modern_vs_tradisional"] = marg_kom.komoditas.map(kom.set_index("nama_dashboard").modern_vs_tradisional_pct)
    R_mk = tabel_data(wb, "d_margin_komoditas", marg_kom, "d_margin_komoditas", {"harga_produsen": RP, "harga_pedagang_besar": RP, "harga_pasar": RP,
                      "margin_pedagang_besar": PCT, "margin_pasar": PCT, "hemat_produsen": PCT, "hemat_pedagang_besar": PCT, "tambahan_pedagang_besar_rp": RP,
                      "tambahan_pasar_rp": RP, "modern_vs_tradisional": "+0%;-0%;0%"})
    tind = T["fact_tindakan"].sort_values(["komoditas_key", "prioritas"])
    utama = tind.groupby("komoditas").head(1).set_index("komoditas").reindex(urut_kom).reset_index()[["komoditas", "tindakan", "alasan_singkat", "nilai", "satuan", "pembanding", "sumber"]]
    R_tu = tabel_data(wb, "d_tindakan", utama, "d_tindakan")
    tabel_data(wb, "d_tindakan_semua", tind[["komoditas", "prioritas", "tindakan", "alasan_singkat", "nilai", "satuan", "pembanding", "sumber"]], "d_tindakan_semua")
    hal = T["fact_hal"].sort_values("urutan")[["urutan", "kalimat", "nilai", "satuan", "pembanding", "sumber"]]
    R_hal = tabel_data(wb, "d_hal", hal, "d_hal")
    w = T["fact_harga_mingguan"]
    wt = w[(w.jenis_pasar_key == 1) & w.kota_key.isin(tujuan)]
    ht = wt.groupby(["komoditas", "minggu_key"]).agg(harga_kota_tujuan=("harga", "median"), harga_nasional=("harga_nasional", "median"),
                                                     perubahan_4m=("perubahan_4m", "median")).reset_index()
    ht["minggu"] = pd.to_datetime(ht.minggu_key.astype(str), format="%Y%m%d")
    ht["selisih_vs_nasional"] = ht.harga_kota_tujuan / ht.harga_nasional - 1
    ht = ht[["komoditas", "minggu_key", "minggu", "harga_kota_tujuan", "harga_nasional", "selisih_vs_nasional", "perubahan_4m"]]
    R_ht = tabel_data(wb, "d_harga_tujuan", ht, "d_harga_tujuan", {"minggu": "yyyy-mm-dd", "harga_kota_tujuan": "#,##0", "harga_nasional": "#,##0", "selisih_vs_nasional": PCT, "perubahan_4m": PCT})
    n_ht = len(ht)

    # ================================================================ RINGKASAN
    ws = S["ringkasan"]
    ht_akhir = ht[ht.minggu_key == akhir].set_index("komoditas")
    tercepat = ht_akhir.perubahan_4m.idxmax()
    beli_awal = int(tind[tind.beli_lebih_awal].komoditas.nunique())
    n_layak, n_aman = int(rute.layak.sum()), int((rute.layak & rute.layak_jika_jual_ke_pedagang_besar).sum())
    teks(ws, "A2", f'=IF(kpi_beli_awal=0,"Minggu ini belum ada komoditas yang perlu dibeli lebih awal; ","Minggu ini "&kpi_beli_awal&" komoditas perlu dibeli lebih awal; ")&kpi_rute_untung&" rute pindah stok masih memberi margin"', 15, True, wrap=True)
    teks(ws, "A3", f"{label_minggu} · harga pasar tradisional di kota tujuan (Jakarta Pusat, Bandung, Semarang, Surabaya) · {info['data_terakhir']} data terakhir", 9, color=MUTED)
    # kartu (baris 5-7): nilai = rumus dari sheet data
    nama(wb, "kpi_beli_awal", f"{q(NAMA_SHEET['ringkasan'])}!$B$6")
    nama(wb, "kpi_rute_untung", f"{q(NAMA_SHEET['ringkasan'])}!$E$6")
    nama(wb, "kpi_kenaikan_tercepat", f"{q(NAMA_SHEET['ringkasan'])}!$H$6")
    nama(wb, "kpi_peringatan_aktif", f"{q(NAMA_SHEET['ringkasan'])}!$K$6")
    kartu(ws, 2, 5, "Perlu dibeli lebih awal (komoditas)", f'=COUNTIF({q("d_tindakan_semua")}!$C$2:$C$200,"Amankan stok*")+COUNTIF({q("d_tindakan_semua")}!$C$2:$C$200,"Beli lebih awal*")', "0",
          sub="Lebaran masih jauh; kenaikan musiman belum mulai" if beli_awal == 0 else "masuk masa menjelang Lebaran")
    kartu(ws, 5, 5, "Rute pindah stok yang untung (dari 32)", f"=COUNTIF({R_rute['layak']},TRUE)", "0", TEAL,
          sub=f"=COUNTIFS({R_rute['layak']},TRUE,{R_rute['layak_jika_jual_ke_pedagang_besar']},TRUE)&\" rute aman di semua skenario (ongkos kirim ×{kons:g} perkiraan)\"")
    kartu(ws, 8, 5, "Kenaikan tercepat (vs 4 minggu lalu)",
          f"=INDEX({R_ht['perubahan_4m']},MATCH(\"{tercepat}|{akhir}\",INDEX({R_ht['komoditas']}&\"|\"&{R_ht['minggu_key']},0),0))", "+0%;-0%;0%", ORANGE_TXT,
          sub=f"{tercepat.lower()}, sekarang Rp {rp(ht_akhir.loc[tercepat, 'harga_kota_tujuan'])}/kg di kota tujuan")
    al_akhir = T["fact_alert"]
    n_peringatan = int(al_akhir[(al_akhir.minggu_key == al_akhir.minggu_key.max()) & al_akhir.peringatan].kota_key.nunique())
    kartu(ws, 11, 5, "Peringatan dini aktif (kota)", n_peringatan, "0", sub="hanya informasi: sistem ini benar 1 dari 10 kali, belum bisa diandalkan")
    # tabel saran
    teks(ws, "A9", "Yang kami sarankan untuk tiap komoditas (tindakan prioritas tertinggi)", 11, True)
    for j, h in enumerate(["Komoditas", "Tindakan", "Alasan singkat"], 1):
        teks(ws, f"{get_column_letter([1, 2, 6][j - 1])}10", h, 9, True, MUTED, bg=WHITE)
    ws.merge_cells("B10:E10"); ws.merge_cells("F10:N10")
    for i in range(6):
        r = 11 + i
        teks(ws, f"A{r}", f"=INDEX({R_tu['komoditas']},{i + 1})", 10, True, bg=WHITE)
        ws.merge_cells(f"B{r}:E{r}"); teks(ws, f"B{r}", f"=INDEX({R_tu['tindakan']},{i + 1})", 10, False, TEAL, wrap=True, bg=WHITE)
        ws.merge_cells(f"F{r}:N{r}"); teks(ws, f"F{r}", f"=INDEX({R_tu['alasan_singkat']},{i + 1})", 9, False, MUTED, wrap=True, bg=WHITE)
        ws.row_dimensions[r].height = 28
    # tiga hal
    teks(ws, "A18", "Tiga hal yang perlu diketahui minggu ini", 11, True)
    for i in range(3):
        r = 19 + i
        ws.merge_cells(f"A{r}:N{r}")
        teks(ws, f"A{r}", f"=INDEX({R_hal['urutan']},{i + 1})&\". \"&INDEX({R_hal['kalimat']},{i + 1})", 10, False, INK, wrap=True, bg=WHITE)
        ws.row_dimensions[r].height = 34
    # grafik harga cabai rawit vs nasional (52 minggu terakhir)
    sub_ht = ht[ht.komoditas == "Cabai rawit merah"].tail(52)
    idx0 = ht.index.get_loc(sub_ht.index[0]) + 2
    idx1 = ht.index.get_loc(sub_ht.index[-1]) + 2
    wsd = wb["d_harga_tujuan"]
    # judul grafik = kesimpulan (selisih minggu ini vs nasional); angkanya di sel tersembunyi X6 untuk audit
    nama(wb, "kpi_selisih_nasional_rawit", f"{q(NAMA_SHEET['ringkasan'])}!$X$6")
    ws["X6"] = f"=INDEX({R_ht['selisih_vs_nasional']},MATCH(\"Cabai rawit merah|{akhir}\",INDEX({R_ht['komoditas']}&\"|\"&{R_ht['minggu_key']},0),0))"
    ws["X6"].number_format = "+0%;-0%;0%"
    ws.column_dimensions["X"].hidden = True
    teks(ws, "A23", '="Harga cabai rawit di kota tujuan "&TEXT(ABS(X6),"0%")&IF(X6<0," di bawah"," di atas")&" rata-rata nasional minggu ini"', 11, True)
    ch = LineChart()
    ch.title = "Cabai rawit: kota tujuan vs rata-rata nasional (Rp/kg), 52 minggu terakhir"
    ch.add_data(Reference(wsd, min_col=4, min_row=idx0 - 1, max_col=5, max_row=idx1), titles_from_data=False)
    ch.set_categories(Reference(wsd, min_col=3, min_row=idx0, max_row=idx1))
    ch.series[0].tx = None; ch.series[1].tx = None
    from openpyxl.chart.series import SeriesLabel
    ch.series[0].tx = SeriesLabel(v="Kota tujuan"); ch.series[1].tx = SeriesLabel(v="Nasional")
    ch.series[0].graphicalProperties.line.solidFill = TEAL; ch.series[0].graphicalProperties.line.width = 22000
    ch.series[1].graphicalProperties.line.solidFill = GRAY; ch.series[1].graphicalProperties.line.width = 16000
    ch.y_axis.numFmt = "#,##0"; sumbu_bulan(ch); ch.height, ch.width = 7.5, 33
    rapikan(ch)
    ws.add_chart(ch, "A24")
    # keputusan dari data minggu ini (bukan teks tetap): amankan stok -> pindah stok (rute aman dulu) -> harga naik cepat
    kep = [(f"{t_.tindakan}: {t_.komoditas.lower()}.", tanpa_awalan(t_.alasan_singkat)) for t_ in tind[tind.beli_lebih_awal].itertuples()]
    pin = tind[tind.pindah_stok].merge(rute[["rute_key", "layak_jika_jual_ke_pedagang_besar", "minggu_untung_berturut_terakhir"]], on="rute_key")
    pin = pin.sort_values(["layak_jika_jual_ke_pedagang_besar", "nilai"], ascending=False)
    for t_ in pin.head(2).itertuples():
        asal, tuju = [x.strip().replace("Kota ", "") for x in t_.tindakan.replace("Pindahkan stok ", "").split("→")]
        k = t_.komoditas.lower()
        kep.append((f"Pindahkan stok {k} dari {asal} ke {tuju} minggu ini.",
                    f"{k} di {asal} Rp {rp(t_.nilai)}/kg lebih murah daripada di {tuju} setelah ongkos kirim, {int(t_.minggu_untung_berturut_terakhir)} minggu berturut-turut; "
                    + ("tetap untung walau dijual ke pedagang besar" if t_.layak_jika_jual_ke_pedagang_besar else "untung hanya jika dijual eceran")))
    naik = ht_akhir.loc[tercepat, "perubahan_4m"]
    if naik >= cfg["insights"]["ambang_naik_4_minggu"]:
        kep.append((f"Beli {tercepat.lower()} secukupnya dulu; tunda pembelian besar.",
                    f"harga {tercepat.lower()} naik {naik:.0%} dalam 4 minggu, sekarang Rp {rp(ht_akhir.loc[tercepat, 'harga_kota_tujuan'])}/kg"))
    else:
        kep.append(("Beli dengan pola rutin; tidak ada harga yang naik cepat.", f"kenaikan 4 minggu tertinggi hanya {naik:.0%} ({tercepat.lower()})"))
    keputusan(ws, 40, kep[:3])
    teks(ws, "A44",f"Ongkos kirim adalah perkiraan kami, bukan tarif aktual · harga di pasar belum tentu harga jual kita · data PIHPS (Bank Indonesia) diambil {info['tanggal_pengambilan']}", 8, color=MUTED)

    # ================================================================ GEJOLAK HARGA
    ws = S["gejolak"]
    GJ = q(NAMA_SHEET["gejolak"])
    teks(ws, "A2", f'=INDEX({R_kom["nama_dashboard"]},MATCH(pilih_komoditas,{R_kom["nama_dashboard"]},0))&" berubah rata-rata "&TEXT(INDEX({R_kom["rata_perubahan_per_minggu"]},MATCH(pilih_komoditas,{R_kom["nama_dashboard"]},0)),"0%")&" setiap minggu, terutama di "&SUBSTITUTE(INDEX(AA12:AA46,1),"Kota ","")', 15, True, wrap=True)
    teks(ws, "A3", "Gejolak = besarnya perubahan harga dari minggu ke minggu di pasar tradisional, Jul 2023 – Okt 2026 · pilih komoditas di sel kuning", 9, color=MUTED)
    teks(ws, "A5", "Komoditas:", 10, True)
    teks(ws, "B5", "Cabai rawit merah", 10, True, INK, bg="FFF3C4")
    ws.merge_cells("B5:D5")
    dv = DataValidation(type="list", formula1=f"={R_kom['nama_dashboard']}", allow_blank=False)
    ws.add_data_validation(dv); dv.add("B5")
    nama(wb, "pilih_komoditas", f"{GJ}!$B$5")
    # helper: gejolak per kota untuk komoditas terpilih (kolom X..AB, baris 12..46), diurutkan dengan LARGE
    kota_list = sorted(g.kota.unique())
    teks(ws, "X10", "Bantu (jangan diubah): kota | gejolak komoditas terpilih | pengurut | peringkat | kota urut | nilai urut", 8, color=MUTED)
    for i, k in enumerate(kota_list):
        r = 12 + i
        ws[f"X{r}"] = k
        ws[f"Y{r}"] = f"=IFERROR(SUMIFS({R_gej['gejolak_mingguan']},{R_gej['kota']},X{r},{R_gej['komoditas']},pilih_komoditas),0)"
        ws[f"Z{r}"] = f"=Y{r}+ROW()/1000000000"
        ws[f"AA{r}"] = f"=INDEX($X$12:$X$46,MATCH(LARGE($Z$12:$Z$46,{i + 1}),$Z$12:$Z$46,0))"
        ws[f"AB{r}"] = f"=INDEX($Y$12:$Y$46,MATCH(LARGE($Z$12:$Z$46,{i + 1}),$Z$12:$Z$46,0))"
        ws[f"AB{r}"].number_format = "0.0%"
    for col in "XYZAA".split():
        pass
    for col in ("X", "Y", "Z", "AA", "AB"):
        ws.column_dimensions[col].hidden = True
    nama(wb, "kpi_kota_paling_tidak_stabil", f"{GJ}!$AA$12")
    nama(wb, "kpi_gejolak_kota_teratas", f"{GJ}!$AB$12")
    # grafik batang 10 kota teratas
    bc = BarChart(); bc.type = "bar"; bc.style = 10
    bc.title = "Kota paling tidak stabil (komoditas terpilih)"
    bc.add_data(Reference(ws, min_col=28, min_row=12, max_row=21), titles_from_data=False)
    bc.set_categories(Reference(ws, min_col=27, min_row=12, max_row=21))
    bc.series[0].graphicalProperties.solidFill = ORANGE; bc.legend = None; bc.y_axis.numFmt = "0%"; bc.x_axis.scaling.orientation = "maxMin"; bc.y_axis.crosses = "max"
    bc.height, bc.width = 7.5, 11
    rapikan(bc)
    ws.add_chart(bc, "H5")
    # peta gelembung: lon (x), lat (y), ukuran = gejolak
    teks(ws, "AD10", "Bantu peta: kota | lon | lat | gejolak terpilih", 8, color=MUTED)
    for i, k in enumerate(kota_list):
        r = 12 + i
        ws[f"AD{r}"] = k
        ws[f"AE{r}"] = f"=INDEX({R_kota['lon']},MATCH(AD{r},{R_kota['kota']},0))"
        ws[f"AF{r}"] = f"=INDEX({R_kota['lat']},MATCH(AD{r},{R_kota['kota']},0))"
        ws[f"AG{r}"] = f"=Y{r}*100"
    for col in ("AD", "AE", "AF", "AG"):
        ws.column_dimensions[col].hidden = True
    bub = BubbleChart(); bub.style = 18
    bub.title = "Letak kota di Pulau Jawa, barat ke timur (besar gelembung = gejolak)"
    s = Series(values=Reference(ws, min_col=32, min_row=12, max_row=46), xvalues=Reference(ws, min_col=31, min_row=12, max_row=46),
               zvalues=Reference(ws, min_col=33, min_row=12, max_row=46), title="Kota")
    s.graphicalProperties.solidFill = ORANGE
    bub.series.append(s); bub.legend = None; bub.height, bub.width = 7.5, 13
    bub.x_axis.scaling.min = 105.5; bub.x_axis.scaling.max = 114.5
    bub.y_axis.scaling.min = -9; bub.y_axis.scaling.max = -5.5; bub.x_axis.tickLblPos = "low"; bub.bubbleScale = 25
    rapikan(bub)
    ws.add_chart(bub, "A7")
    # heatmap kota x komoditas (SUMIFS) + color scale
    teks(ws, "A24", "Cabai tidak stabil di semua kota; ayam, telur, dan bawang putih stabil di mana pun (semua komoditas; tidak ikut pilihan)", 11, True)
    teks(ws, "A25", "Kota", 9, True, MUTED, bg=WHITE)
    for j, k in enumerate(urut_kom):
        teks(ws, f"{get_column_letter(2 + j)}25", k, 9, True, MUTED, bg=WHITE, wrap=True)
    ws.row_dimensions[25].height = 28
    top_kota = (g.pivot(index="kota", columns="komoditas", values="gejolak_mingguan")["Cabai rawit merah"].sort_values(ascending=False).index.tolist())
    for i, k in enumerate(top_kota):
        r = 26 + i
        teks(ws, f"A{r}", k, 9, bg=WHITE)
        for j, kk in enumerate(urut_kom):
            kol = f"{get_column_letter(2 + j)}$25"   # kota tanpa data komoditas ini -> "–", bukan 0%
            c = ws.cell(r, 2 + j, f'=IF(COUNTIFS({R_gej["kota"]},$A{r},{R_gej["komoditas"]},{kol})=0,"–",SUMIFS({R_gej["gejolak_mingguan"]},{R_gej["kota"]},$A{r},{R_gej["komoditas"]},{kol}))')
            c.number_format = "0%"; c.font = font(9, color=INK); c.alignment = Alignment(horizontal="center")
    ws.conditional_formatting.add(f"B26:G{25 + len(top_kota)}", ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF", end_type="num", end_value=0.25, end_color=ORANGE))
    keputusan(ws, 27 + len(top_kota), [
        ("Pantau harga kedua cabai setiap minggu dan siapkan stok cadangan.", "harga cabai rawit dan keriting berubah 13% dan 10% per minggu, lima kali lebih besar daripada telur"),
        ("Cukup dipantau bulanan; beli dengan pola rutin.", "ayam, telur, dan bawang putih hanya berubah 2–3% per minggu di semua kota"),
        ("Verifikasi harga Blitar di lapangan sebelum dipakai untuk keputusan.", "Blitar masuk tiga kota paling tidak stabil di 5 dari 6 komoditas, tetapi datanya hanya dari 2 pasar")])

    # ================================================================ PINDAH STOK
    ws = S["rute"]
    PS = q(NAMA_SHEET["rute"])
    teks(ws, "A2", '=kpi_rute_untung_dinamis&" dari 32 rute pindah stok masih untung setelah ongkos kirim; margin tertinggi "&LOWER(LEFT($AA$12,1))&MID($AA$12,2,200)&", Rp "&'
         + rupiah_teks("$AD$12") + '&"/kg"', 15, True, wrap=True)
    teks(ws, "A3", "Harga beli = pedagang besar di kota asal; harga jual = pasar tradisional di kota tujuan · median mingguan Jul 2023 – Okt 2026 · ubah sel kuning untuk skenario ongkos", 9, color=MUTED)
    teks(ws, "A5", "Ongkos kirim ×", 10, True); teks(ws, "B5", kons, 10, True, bg="FFF3C4", fmt="0.00")
    dvp = DataValidation(type="list", formula1='"0.5,1,1.5,2"', allow_blank=False); ws.add_data_validation(dvp); dvp.add("B5")
    teks(ws, "C5", "Tambahan ongkos (Rp/kg)", 10, True); teks(ws, "E5", 0, 10, True, bg="FFF3C4", fmt="#,##0")
    teks(ws, "F5", "← perkiraan kami, bukan tarif aktual; 1,5× = skenario konservatif", 8, color=MUTED)
    nama(wb, "pengali_ongkos", f"{PS}!$B$5"); nama(wb, "tambahan_ongkos", f"{PS}!$E$5")
    # helper per rute (kolom X..AF baris 12..43): key | nama | selisih median | margin (pengali, tambahan) | minggu untung dari 10 | risiko | pengurut | urutan
    teks(ws, "X10", "Bantu (jangan diubah)", 8, color=MUTED)
    n_r = len(d_rute)
    for i in range(n_r):
        r = 12 + i
        ws[f"X{r}"] = f"=INDEX({R_rute['rute_key']},{i + 1})"
        ws[f"Y{r}"] = f"=INDEX({R_rute['nama_rute']},{i + 1})"
        ws[f"Z{r}"] = f"=INDEX({R_rute['median_selisih_harga_rp_per_kg']},{i + 1})"
        # margin median pada pengali terpilih (dari tabel sensitivitas) dikurangi tambahan ongkos x pengali (pergeseran konstan -> median ikut bergeser)
        ws[f"AB{r}"] = f"=INDEX({R_sens['median_untung_rp_per_kg']},MATCH(X{r}&\"|\"&pengali_ongkos,INDEX({R_sens['rute_key']}&\"|\"&{R_sens['pengali_ongkos']},0),0))-tambahan_ongkos*pengali_ongkos"
        ws[f"AC{r}"] = (f"=10*SUMPRODUCT(({R_sp['rute_key']}=X{r})*ISNUMBER({R_sp['selisih_harga']})*(({R_sp['selisih_harga']}-({R_sp['ongkos_angkut']}+{R_sp['susut_rp']}+tambahan_ongkos)*pengali_ongkos)>0))"
                        f"/SUMPRODUCT(({R_sp['rute_key']}=X{r})*ISNUMBER({R_sp['selisih_harga']}))")
        ws[f"AE{r}"] = f"=INDEX({R_rute['risiko']},{i + 1})"
        ws[f"AF{r}"] = f"=AB{r}+ROW()/1000000"
        ws[f"AA{r}"] = f"=INDEX($Y$12:$Y$43,MATCH(LARGE($AF$12:$AF$43,{i + 1}),$AF$12:$AF$43,0))"
        ws[f"AD{r}"] = f"=INDEX($AB$12:$AB$43,MATCH(LARGE($AF$12:$AF$43,{i + 1}),$AF$12:$AF$43,0))"
        ws[f"AG{r}"] = f"=INDEX($X$12:$X$43,MATCH(LARGE($AF$12:$AF$43,{i + 1}),$AF$12:$AF$43,0))"
    for col in ("X", "Y", "Z", "AA", "AB", "AC", "AD", "AE", "AF", "AG"):
        ws.column_dimensions[col].hidden = True
    nama(wb, "kpi_rute_untung_dinamis", f"{PS}!$B$8")
    nama(wb, "kpi_rute_aman", f"{PS}!$E$8")
    nama(wb, "kpi_rute_margin_positif", f"{PS}!$H$8")
    nama(wb, "kpi_margin_rute_terbaik", f"{PS}!$K$8")
    kartu(ws, 2, 7, "Rute yang untung (dari 32)", "=SUMPRODUCT(($AB$12:$AB$43>0)*($AC$12:$AC$43>=6))", "0", TEAL, sub="untung minimal 6 dari 10 minggu dan margin median > 0, pada skenario di sel kuning")
    kartu(ws, 5, 7, "Rute aman di semua skenario", f"=COUNTIFS({R_rute['layak']},TRUE,{R_rute['layak_jika_jual_ke_pedagang_besar']},TRUE)", "0", TEAL, sub="tetap untung walau dijual ke pedagang besar (dihitung pada ×1,5)")
    kartu(ws, 8, 7, "Rute dengan margin positif", "=COUNTIF($AB$12:$AB$43,\">0\")", "0", sub="median margin > 0 pada skenario di sel kuning")
    kartu(ws, 11, 7, "Margin rute terbaik (Rp/kg)", "=INDEX($AD$12:$AD$43,1)", RP, TEAL, sub='=INDEX($AA$12:$AA$43,1)')
    # tabel rute terurut
    teks(ws, "A11", "Rute diurutkan dari margin terbesar per kg (ongkos kirim sesuai skenario)", 11, True)
    hdr = ["Rute", "", "", "Selisih harga", "Ongkos kirim", "Margin/kg", "Minggu untung (dari 10)", "Risiko"]
    for j, h in enumerate(hdr, 1):
        if h:
            teks(ws, f"{get_column_letter(j)}12", h, 9, True, MUTED, bg=WHITE, wrap=True)
    ws.merge_cells("A12:C12")
    ws.row_dimensions[12].height = 28
    ws.column_dimensions["G"].width = 16
    for i in range(n_r):
        r = 13 + i
        ws.merge_cells(f"A{r}:C{r}")
        teks(ws, f"A{r}", f"=INDEX($AA$12:$AA$43,{i + 1})", 9, bg=WHITE)
        teks(ws, f"D{r}", f"=INDEX($Z$12:$Z$43,MATCH(INDEX($AG$12:$AG$43,{i + 1}),$X$12:$X$43,0))", 9, bg=WHITE, fmt="#,##0", align="right")
        teks(ws, f"E{r}", f"=D{r}-F{r}", 9, bg=WHITE, fmt="#,##0", align="right")
        teks(ws, f"F{r}", f"=INDEX($AD$12:$AD$43,{i + 1})", 9, True, TEAL, bg=WHITE, fmt="#,##0", align="right")
        teks(ws, f"G{r}", f"=INDEX($AC$12:$AC$43,MATCH(INDEX($AG$12:$AG$43,{i + 1}),$X$12:$X$43,0))", 9, bg=WHITE, fmt="0.0", align="right")
        teks(ws, f"H{r}", f"=INDEX($AE$12:$AE$43,MATCH(INDEX($AG$12:$AG$43,{i + 1}),$X$12:$X$43,0))", 9, bg=WHITE)
    # grafik rute terbaik (kunci rute terbaik dinamis): 52 minggu terakhir
    minggu_52 = sorted(sp.minggu_key.unique())[-52:]
    teks(ws, "AI10", "Bantu grafik: minggu | selisih | ongkos (rute terbaik)", 8, color=MUTED)
    for i, mk in enumerate(minggu_52):
        r = 12 + i
        ws[f"AI{r}"] = pd.to_datetime(str(mk), format="%Y%m%d").to_pydatetime(); ws[f"AI{r}"].number_format = BULAN_SUMBU
        key = f"INDEX($AG$12:$AG$43,1)&\"|{mk}\""
        ws[f"AJ{r}"] = f"=IFERROR(INDEX({R_sp['selisih_harga']},MATCH({key},{R_sp['kunci']},0)),NA())"
        ws[f"AK{r}"] = f"=IFERROR((INDEX({R_sp['ongkos_angkut']},MATCH({key},{R_sp['kunci']},0))+INDEX({R_sp['susut_rp']},MATCH({key},{R_sp['kunci']},0))+tambahan_ongkos)*pengali_ongkos,NA())"
    for col in ("AI", "AJ", "AK"):
        ws.column_dimensions[col].hidden = True
    lc = LineChart(); lc.title = "Rute terbaik: selisih harga vs ongkos kirim (Rp/kg), 52 minggu terakhir"
    lc.add_data(Reference(ws, min_col=36, min_row=12, max_col=37, max_row=63), titles_from_data=False)
    lc.set_categories(Reference(ws, min_col=35, min_row=12, max_row=63))
    from openpyxl.chart.series import SeriesLabel as SL
    lc.series[0].tx = SL(v="Selisih harga"); lc.series[1].tx = SL(v="Ongkos kirim (skenario)")
    lc.series[0].graphicalProperties.line.solidFill = TEAL; lc.series[1].graphicalProperties.line.solidFill = ORANGE
    lc.series[1].graphicalProperties.line.dashStyle = "dash"; lc.y_axis.numFmt = "#,##0"; sumbu_bulan(lc, tiap=2); lc.height, lc.width = 8, 13
    rapikan(lc)
    ws.add_chart(lc, "J11")
    keputusan(ws, 46, [
        ("Jalankan 9 rute aman lebih dulu; risikonya paling kecil.", "9 rute tetap untung walau dijual ke pedagang besar: cabai rawit dari Kediri dan Jember, bawang merah dari Tegal dan Probolinggo"),
        ("Gunakan rute telur, ayam, dan cabai keriting hanya untuk pasokan ke gerai ritel sendiri.", "rute itu hanya untung jika dijual eceran; tidak ada yang bertahan jika dijual ke pedagang besar"),
        ("Tunda rute yang marginnya di bawah Rp 1.000/kg.", "jika ongkos kirim dua kali perkiraan, rute yang untung turun dari 27 menjadi 22")])

    # ================================================================ PERINGATAN DINI
    ws = S["peringatan"]
    PD = q(NAMA_SHEET["peringatan"])
    teks(ws, "A2", '=IF(kpi_benar_dari_10<kpi_benar_tahun_lalu,"Sistem peringatan dini belum bisa diandalkan: hanya "&TEXT(kpi_benar_dari_10,"0")&" dari 10 peringatan yang terbukti benar","Peringatan dini terbukti benar "&SUBSTITUTE(FIXED(kpi_benar_dari_10,1),".",",")&" dari 10 kali, lebih baik daripada pola tahun lalu")', 15, True, wrap=True)
    teks(ws, "A3", "Diuji pada data 2025–2026, pasar tradisional; setiap tahun uji hanya memakai pengetahuan dari data sebelumnya", 9, color=MUTED)
    nama(wb, "kpi_benar_dari_10", f"{PD}!$B$6"); nama(wb, "kpi_tertangkap_dari_10", f"{PD}!$E$6")
    nama(wb, "kpi_lebih_awal", f"{PD}!$H$6"); nama(wb, "kpi_benar_tahun_lalu", f"{PD}!$K$6")
    kartu(ws, 2, 5, "Peringatan yang terbukti benar (dari 10)", f"=INDEX({R_uji['benar_dari_10']},1)", "0.0", ORANGE_TXT, sub="lebih buruk daripada memberi peringatan setiap minggu (1,6)")
    kartu(ws, 5, 5, "Lonjakan yang terdeteksi (dari 10)", f"=INDEX({R_uji['tertangkap_dari_10']},1)", "0.0", sub=f"dari {int(uji.jumlah_lonjakan.iloc[0])} lonjakan harga")
    kartu(ws, 8, 5, "Datang lebih awal (minggu)", f"=INDEX({R_uji['lebih_awal_minggu']},1)", "0.0", sub="tidak berarti sendirian; tebakan asal pun 7,2 minggu")
    kartu(ws, 11, 5, "Pola tahun lalu: benar (dari 10)", f"=INDEX({R_uji['benar_dari_10']},4)", "0.0", TEAL, sub="mendeteksi 4,2 dari 10 lonjakan — lebih baik daripada sistem")
    # bar aturan
    teks(ws, "A9", "Peringatan yang benar dari 10: semua versi aturan kalah dari pola tahun lalu", 11, True)
    for i in range(5):
        r = 10 + i
        teks(ws, f"A{r}", f"=INDEX({R_uji['aturan']},{i + 1})", 9, bg=WHITE); ws.merge_cells(f"A{r}:C{r}")
        teks(ws, f"D{r}", f"=INDEX({R_uji['benar_dari_10']},{i + 1})", 9, True, bg=WHITE, fmt="0.0", align="right")
    bc2 = BarChart(); bc2.type = "bar"; bc2.title = "Peringatan yang benar (dari 10) per aturan"
    bc2.add_data(Reference(ws, min_col=4, min_row=10, max_row=14), titles_from_data=False); bc2.set_categories(Reference(ws, min_col=1, min_row=10, max_row=14))
    bc2.series[0].graphicalProperties.solidFill = ORANGE; bc2.legend = None; bc2.height, bc2.width = 6, 14; bc2.x_axis.scaling.orientation = "maxMin"; bc2.y_axis.crosses = "max"
    rapikan(bc2)
    ws.add_chart(bc2, "A16")
    # tabel kota naik lebih dulu
    teks(ws, "H9", "Kota yang harganya naik lebih dulu (selisih 1 minggu; 3 teratas per komoditas)", 11, True)
    # komoditas H:I, kota J:K (digabung agar nama panjang tidak terpotong), jeda L, kekuatan M
    for r in range(10, 11 + len(pa)):
        ws.merge_cells(f"H{r}:I{r}"); ws.merge_cells(f"J{r}:K{r}")
    for col, h in zip("HJLM", ["Komoditas", "Kota", "Jeda (minggu)", "Hubungan"]):
        teks(ws, f"{col}10", h, 9, True, MUTED, bg=WHITE)
    for i in range(len(pa)):
        r = 11 + i
        for col, c in zip("HJLM", ["komoditas", "kota", "jeda_minggu", "hubungan"]):
            teks(ws, f"{col}{r}", f"=INDEX({R_pa[c]},{i + 1})", 9, bg=WHITE, align="center" if col == "L" else "left")
    # timeline Jakarta
    teks(ws, "A30", "Contoh Jakarta Pusat: harga cabai rawit (garis) dan minggu dengan peringatan (titik); tak satu pun tepat sebelum lonjakan", 11, True)
    wsa = wb["d_peringatan_jakarta"]
    tl = LineChart(); tl.title = "Cabai rawit Jakarta Pusat 2025–2026 (Rp/kg)"
    tl.add_data(Reference(wsa, min_col=2, min_row=1, max_col=3, max_row=len(al) + 1), titles_from_data=True)
    tl.set_categories(Reference(wsa, min_col=1, min_row=2, max_row=len(al) + 1))
    from openpyxl.chart.series import SeriesLabel as SL4
    tl.series[0].tx = SL4(v="Harga"); tl.series[1].tx = SL4(v="Minggu dengan peringatan")
    tl.series[0].graphicalProperties.line.solidFill = TEAL
    tl.series[1].graphicalProperties.line.noFill = True; tl.series[1].marker.symbol = "circle"; tl.series[1].marker.size = 8
    tl.series[1].marker.graphicalProperties.solidFill = ORANGE; tl.series[1].marker.graphicalProperties.line.solidFill = ORANGE
    tl.y_axis.numFmt = "#,##0"; sumbu_bulan(tl, tiap=2); tl.height, tl.width = 7, 30
    rapikan(tl)
    ws.add_chart(tl, "A31")
    keputusan(ws, 47, [
        ("Jangan gunakan peringatan dini sebagai dasar pembelian.", "peringatan sistem hanya benar 1,2 dari 10 kali; pola sederhana \"sama seperti tahun lalu\" benar 2,2 dari 10"),
        ("Jadikan kalender musiman sebagai pengingat jadwal beli.", "pola musiman lebih dapat dipercaya: bawang merah naik 4 minggu dan ayam 6 minggu sebelum Lebaran"),
        ("Pantau kota yang naik lebih dulu sebagai informasi tambahan, bukan pemicu beli.", "kota itu hanya mendahului 1 minggu, dan untuk kedua cabai hubungannya hanya sedang")])

    # ================================================================ POLA MUSIMAN
    ws = S["pola"]
    bm = kom.set_index("nama_dashboard").loc["Bawang merah"]
    teks(ws, "A2", f'="Bawang merah naik "&TEXT(INDEX({R_kom["idulfitri_kenaikan_minggu_min1"]},MATCH("Bawang merah",{R_kom["nama_dashboard"]},0)),"0%")&" menjelang Lebaran dan mulai naik "&ABS(INDEX({R_kom["idulfitri_mulai_naik_minggu"]},MATCH("Bawang merah",{R_kom["nama_dashboard"]},0)))&" minggu sebelumnya; stok harus masuk sebelum itu"', 15, True, wrap=True)
    teks(ws, "A3", "Kenaikan vs harga normal (rata-rata 9–12 minggu sebelum Lebaran), median antar kota, pasar tradisional · Lebaran 2024, 2025, 2026 · garis = rata-rata, pita = rentang antar tahun", 9, color=MUTED)
    kol_mulai = R_kom["idulfitri_mulai_naik_minggu"].split("!$")[1].split("$")[0]   # huruf kolom di sheet data
    nama(wb, "kpi_bawang_merah_mulai_naik", f"{q('d_komoditas')}!${kol_mulai}${2 + urut_kom.index('Bawang merah')}")
    wsp = wb["d_pola_lebaran"]
    for i, k in enumerate(urut_kom):
        r0 = 2 + i * 13
        c = LineChart()
        row_k = d_kom.set_index("nama_dashboard").loc[k]
        mulai = row_k.idulfitri_mulai_naik_minggu
        tag = f"naik {abs(int(mulai))} minggu sebelum · {row_k.idulfitri_tahun_konsisten} tahun" if pd.notna(mulai) else f"tidak konsisten · {row_k.idulfitri_tahun_konsisten} tahun"
        c.title = f"{k}: {tag}"
        c.add_data(Reference(wsp, min_col=3, min_row=r0 - 1, max_col=5, max_row=r0 + 12), titles_from_data=False)
        c.set_categories(Reference(wsp, min_col=2, min_row=r0, max_row=r0 + 12))
        from openpyxl.chart.series import SeriesLabel as SL2
        for s_, lbl, colr, w_ in zip(c.series, ["rata-rata 3 tahun", "terendah", "tertinggi"], [TEAL, GRAY, GRAY], [28000, 12000, 12000]):
            s_.tx = SL2(v=lbl); s_.graphicalProperties.line.solidFill = colr; s_.graphicalProperties.line.width = w_
        c.y_axis.numFmt = "0%"; c.height, c.width = 6.5, 9; c.legend = None; c.x_axis.tickLblPos = "low"
        rapikan(c, judul_pt=9)
        ws.add_chart(c, f"{'AFK'[i % 3]}{5 + (i // 3) * 14}")
    keputusan(ws, 33, [
        ("Amankan stok bawang merah paling lambat 4 minggu sebelum Lebaran.", "bawang merah naik 26% seminggu sebelum Lebaran, 3 tahun berturut-turut; kenaikan mulai 4 minggu sebelumnya"),
        ("Beli ayam 6 minggu sebelum; beli cabai rawit bertahap, jangan sekaligus.", "ayam naik 6 minggu sebelum (2 dari 3 tahun); cabai rawit sudah tinggi sejak 8 minggu tetapi tiap tahun berbeda jauh"),
        ("Tidak perlu stok lebih awal untuk cabai keriting, telur, dan bawang putih.", "ketiganya tidak menunjukkan kenaikan yang konsisten menjelang Lebaran")])
    teks(ws, "A37", "Baru 3 Lebaran dalam data; hasil masih perkiraan kasar. Minggu Lebaran 2024 dan 2025 tanpa pencatatan (libur).", 8, color=MUTED)

    # ================================================================ RANTAI PASOK
    ws = S["rantai"]
    RT = q(NAMA_SHEET["rantai"])
    teks(ws, "A2", f'="Membeli langsung dari produsen menghemat sampai "&TEXT(MAX({R_mk["hemat_produsen"]}),"0%")&" untuk "&LOWER(INDEX({R_mk["komoditas"]},MATCH(MAX({R_mk["hemat_produsen"]}),{R_mk["hemat_produsen"]},0)))&"; untuk "&LOWER(INDEX({R_mk["komoditas"]},MATCH(MIN({R_mk["hemat_produsen"]}),{R_mk["hemat_produsen"]},0)))&" hanya "&TEXT(MIN({R_mk["hemat_produsen"]}),"0%")', 15, True, wrap=True)
    teks(ws, "A3", "Median bulanan Jul 2023 – Okt 2026, kota dengan ketiga tingkat (6–14 kota per komoditas) · harga produsen jarang diperbarui, angka ini perkiraan", 9, color=MUTED)
    nama(wb, "kpi_hemat_produsen_maks", f"{RT}!$X$6")   # sel bantu tersembunyi (angkanya tampil di judul & tabel)
    ws["X6"] = f"=MAX({R_mk['hemat_produsen']})"
    ws["X6"].number_format = PCT
    ws.column_dimensions["X"].hidden = True
    teks(ws, "A5", "Harga per tingkat dan margin per komoditas (Rp/kg, median antar kota)", 11, True)
    hdr = ["Komoditas", "Produsen", "Pedagang besar", "Pasar", "Margin ped. besar", "Margin pasar", "Hemat beli di produsen", "Hemat lewat ped. besar", "Modern vs tradisional"]
    for j, h in enumerate(hdr, 1):
        teks(ws, f"{get_column_letter(j)}6", h, 9, True, MUTED, bg=WHITE, wrap=True)
    ws.row_dimensions[6].height = 28
    cols = ["komoditas", "harga_produsen", "harga_pedagang_besar", "harga_pasar", "margin_pedagang_besar", "margin_pasar", "hemat_produsen", "hemat_pedagang_besar", "modern_vs_tradisional"]
    fmts = [None, "#,##0", "#,##0", "#,##0", PCT, PCT, PCT, PCT, "+0%;-0%;0%"]
    for i in range(len(marg_kom)):
        r = 7 + i
        for j, (c, f) in enumerate(zip(cols, fmts), 1):
            teks(ws, f"{get_column_letter(j)}{r}", f"=INDEX({R_mk[c]},{i + 1})", 9, j == 7, TEAL if j == 7 else INK, bg=WHITE, fmt=f, align="left" if j == 1 else "right")
    # stacked bar
    sb = BarChart(); sb.type = "bar"; sb.grouping = "stacked"; sb.overlap = 100
    sb.title = "Kenaikan harga di setiap tingkat: produsen → pedagang besar → pasar (Rp/kg)"
    wsm = wb["d_margin_komoditas"]
    sb.add_data(Reference(wsm, min_col=3, min_row=1, max_col=3, max_row=len(marg_kom) + 1), titles_from_data=True)
    sb.add_data(Reference(wsm, min_col=10, min_row=1, max_col=11, max_row=len(marg_kom) + 1), titles_from_data=True)
    sb.set_categories(Reference(wsm, min_col=1, min_row=2, max_row=len(marg_kom) + 1))
    from openpyxl.chart.series import SeriesLabel as SL3
    for s_, colr, lbl in zip(sb.series, [GRAY, BLUE, ORANGE], ["Harga produsen", "Tambahan di pedagang besar", "Tambahan di pasar"]):
        s_.graphicalProperties.solidFill = colr
        s_.tx = SL3(v=lbl)
    sb.x_axis.scaling.orientation = "maxMin"; sb.y_axis.crosses = "max"; sb.y_axis.numFmt = "#,##0"; sb.height, sb.width = 7, 14; sb.legend.position = "b"
    rapikan(sb)
    ws.add_chart(sb, "A13")
    # per kota cabai keriting
    teks(ws, "J5", "Cabai keriting: hemat beli di produsen per kota", 11, True)
    set_widths(ws, {"J": 17, "K": 17})
    mk_ = m[m.komoditas == "Cabai merah keriting"].sort_values("hemat_beli_produsen_pct", ascending=False)
    for j, h in enumerate(["Kota", "Produsen dari", "Produsen", "Pasar", "Hemat"], 10):
        teks(ws, f"{get_column_letter(j)}6", h, 9, True, MUTED, bg=WHITE)
    off = m.index.get_loc(mk_.index[0])
    for i in range(len(mk_)):
        r = 7 + i
        idx = m.index.get_loc(mk_.index[i]) + 1
        for j, (c, f) in enumerate(zip(["kota", "harga_produsen_dari", "harga_produsen", "harga_pasar", "hemat_beli_produsen_pct"], [None, None, "#,##0", "#,##0", PCT]), 10):
            teks(ws, f"{get_column_letter(j)}{r}", f"=INDEX({R_m[c]},{idx})", 9, j == 14, TEAL if j == 14 else INK, bg=WHITE, fmt=f, align="left" if j < 12 else "right")
    for r in range(6, 8 + len(mk_)):   # jarak dari kolom "Modern vs tradisional" di sebelah kiri
        ws[f"J{r}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    keputusan(ws, 30, [
        ("Jajaki pembelian langsung dari produsen untuk cabai keriting, cabai rawit, dan ayam.", "harga produsen 41%, 32%, dan 21% lebih murah daripada pasar; pedagang besar mengambil 15–29%"),
        ("Tetap beli bawang merah dan telur lewat pedagang besar.", "untuk keduanya pedagang besar hanya 9% dan 3% di atas harga produsen"),
        ("Keluarkan pasar modern dari daftar sumber pembelian.", "pasar modern 57–82% lebih mahal daripada pasar tradisional untuk cabai dan bawang merah")])
    teks(ws, "A34", "Harga produsen = tingkat petani/peternak, belum termasuk biaya mengumpulkan dari banyak petani · bawang putih tidak punya data produsen (impor) · "
         "'Produsen dari' = kabupaten sekitar bila kota itu tidak punya data produsen (perlu diverifikasi)", 8, color=MUTED)

    # ================================================================ CATATAN METODE
    ws = ws_met
    ws.sheet_view.showGridLines = False
    set_widths(ws, {"A": 34, "B": 110})
    baris = [
        ("Sumber data", f"PIHPS Nasional, Bank Indonesia (bi.go.id/hargapangan); pasar tradisional, modern, pedagang besar, produsen; Pulau Jawa; {info['data_pertama']} s/d {info['data_terakhir']}; diambil {info['tanggal_pengambilan']}. Data bisa direvisi BI."),
        ("Cara membaca angka", "Judul tiap sheet = kesimpulan. Tiap angka punya pembanding (vs 4 minggu lalu, vs nasional, vs ongkos kirim, dari 10 minggu). Sel kuning = input yang boleh diubah."),
        ("Ongkos kirim", "PERKIRAAN, bukan data: tarif truk CDD + tol + bongkar muat + susut (cabai 5%, bawang merah 3%, telur 2%, ayam 1%). Rute disebut untung hanya bila tetap untung pada ongkos 1,5× perkiraan."),
        ("Gejolak harga", "Simpangan baku perubahan harga mingguan (log-return) per kota × komoditas di pasar tradisional. 'Perubahan per minggu' = rata-rata |perubahan %| minggu ke minggu."),
        ("Pindah stok", "Selisih harga = harga pasar tradisional kota tujuan − harga pedagang besar kota sumber (median mingguan). Margin = selisih − (ongkos + susut + tambahan) × pengali. Rute aman = juga untung bila dijual ke pedagang besar di kota tujuan."),
        ("Kota yang naik lebih dulu", "Korelasi silang kejutan harga mingguan (setelah momentum dibuang) kota vs rata-rata kota lain, jeda 1–4 minggu, koreksi Benjamini–Hochberg. Ini kaitan statistik, bukan sebab-akibat."),
        ("Peringatan dini", "Peringatan = kenaikan mingguan jauh di atas kebiasaan (z > 2) di kota yang naik lebih dulu, atau masuk masa sebelum Lebaran. Lonjakan = naik ≥ persentil 90 kenaikan 4 minggu. Diuji 2025 & 2026 tanpa melihat tahun uji. Benar 1,2 dari 10; pola tahun lalu 2,2 dari 10 → belum layak operasional."),
        ("Pola Lebaran", "Harga minggu −8…+4 relatif terhadap harga normal (rata-rata minggu −12…−9), median antar kota, 3 Lebaran (2024–2026). 'Mulai naik' = setiap minggu sampai −1 ≥ 5% di atas normal pada ≥ 2 dari 3 tahun."),
        ("Rantai pasok", "Median bulanan produsen → pedagang besar → pasar tradisional, kota dengan ketiga tingkat; sebagian kota memakai produsen di kabupaten sekitar (PERLU_VERIFIKASI). Harga produsen jarang diperbarui."),
        ("Batasan", "Sampel PIHPS kecil (2 pasar × 2 pedagang per kota); harga eceran ≠ harga jual distributor; volume & stok tidak diketahui; baru 3 Lebaran; kota sumber dipilih dari pengetahuan umum (PERLU_VERIFIKASI)."),
        ("Reproduksi", "Semua angka berasal dari skrip Python (src/) → data/processed/ → sheet d_* di workbook ini (python -m src.build_excel). Audit: python -m src.audit_dashboard --xlsx."),
        ("Nama sel untuk audit", "kpi_beli_awal, kpi_rute_untung, kpi_kenaikan_tercepat, kpi_selisih_nasional_rawit, pilih_komoditas, kpi_kota_paling_tidak_stabil, kpi_gejolak_kota_teratas, pengali_ongkos, tambahan_ongkos, kpi_rute_untung_dinamis, kpi_rute_aman, kpi_rute_margin_positif, kpi_margin_rute_terbaik, kpi_benar_dari_10, kpi_tertangkap_dari_10, kpi_lebih_awal, kpi_benar_tahun_lalu, kpi_bawang_merah_mulai_naik, kpi_hemat_produsen_maks."),
    ]
    teks(ws, "A1", "Catatan metode, sumber, dan batasan", 14, True)
    for i, (k, v) in enumerate(baris, 3):
        teks(ws, f"A{i}", k, 10, True); teks(ws, f"B{i}", v, 10, wrap=True)
        ws.row_dimensions[i].height = 42 if len(v) < 300 else 72
    teks(ws, "A17", f"Dibuat {datetime.now():%Y-%m-%d %H:%M} oleh python -m src.build_excel", 8, color=MUTED)

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def main() -> None:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--out", default=str(OUT))
    a = p.parse_args()
    out = build(Path(a.out))
    print(f"Dashboard ditulis: {out} ({out.stat().st_size / 1e6:.1f} MB)")


if __name__ == "__main__":
    main()
