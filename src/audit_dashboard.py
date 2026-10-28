"""Tahap 6 — Audit dashboard Excel terhadap data/processed/.

1. NILAI ACUAN: angka yang harus tampil di tiap kartu/judul/keputusan, dihitung langsung dari data/processed/ (tanpa Excel).
2. CEK FILE XLSX (--xlsx): buka workbook (sudah dihitung ulang oleh src/recalc_excel.py), baca nilai sel bernama (kpi_*),
   bandingkan dengan nilai acuan, cari sel rumus bererror, dan cari istilah teknis terlarang (glosarium §6.3) di judul/label.
Keluaran: docs/audit_dashboard_report.md.

    python -m src.audit_dashboard                                   # nilai acuan saja
    python -m src.audit_dashboard --xlsx excel/Dashboard_pangan.xlsx
"""
from __future__ import annotations

import argparse
import re
from datetime import datetime
from pathlib import Path

import pandas as pd

from src.analysis.common import ROOT, load_cfg

PROCESSED = ROOT / "data" / "processed"
ISTILAH_TERLARANG = ["spread", "z-score", "zscore", "granger", " lag", "precision", "recall", "lead time", "fact_", "dim_",
                     "log-return", "log return", "cv ", "koefisien variasi", "benjamini", "q-value", "backtest", "baseline"]
SHEET_DASHBOARD = ["Ringkasan", "Gejolak harga", "Pindah stok", "Peringatan dini", "Pola musiman", "Rantai pasok"]


# ---------------------------------------------------------------- bagian 1: nilai acuan
def _t(nama: str) -> pd.DataFrame:
    return pd.read_parquet(PROCESSED / f"{nama}.parquet")


def nilai_acuan(cfg: dict, lcfg: dict) -> dict:
    """Angka yang harus tampil di dashboard, dihitung ulang dari data/processed/ (bukan dari reports/)."""
    kons = lcfg["sensitivitas"]["skenario_konservatif"]
    q2 = cfg["analisis"]["q2_rute"]
    tujuan = list(cfg["wilayah"]["kota_tujuan"])
    kom = _t("dim_komoditas").set_index("komoditas_key")
    kota = _t("dim_kota").set_index("kota_key").kota
    w = _t("fact_harga_mingguan")
    sp = _t("fact_spread")
    dr = _t("dim_rute")
    akhir = int(w.minggu_key.max())

    t = w[(w.minggu_key == akhir) & (w.jenis_pasar_key == 1) & w.kota_key.isin(tujuan)]
    harga = t.groupby("komoditas_key").agg(harga=("harga", "median"), perubahan_4m=("perubahan_4m", "median"),
                                           harga_nasional=("harga_nasional", "median"))
    harga["nama"] = harga.index.map(kom.nama_dashboard)
    harga["selisih_nasional_pct"] = harga.harga / harga.harga_nasional - 1
    tercepat = harga.sort_values("perubahan_4m", ascending=False).iloc[0]
    tind = _t("fact_tindakan")
    tind = tind[tind.minggu_key == akhir]
    hal = _t("fact_hal")
    al = _t("fact_alert")

    def rute_per_pengali(p: float, tambahan: float = 0.0) -> dict:
        s = sp[sp.selisih_harga.notna()].copy()
        s["untung"] = s.selisih_harga - (s.ongkos_angkut + s.susut_rp + tambahan) * p
        g = s.groupby("rute_key").untung.agg(median="median", porsi=lambda x: (x > 0).mean())
        return {"median_positif": int((g["median"] > 0).sum()),
                "layak": int(((g["porsi"] >= q2["layak_min_porsi_minggu_untung"]) & (g["median"] > 0)).sum()),
                "margin_terbaik": float(g["median"].max())}

    sens = {p: rute_per_pengali(p) for p in lcfg["sensitivitas"]["pengali"]}
    s_akhir = sp[sp.minggu_key == akhir]
    untung_akhir = s_akhir.selisih_harga - (s_akhir.ongkos_angkut + s_akhir.susut_rp) * kons
    terbaik = dr.sort_values("median_untung_rp_per_kg", ascending=False).iloc[0]

    g = _t("fact_gejolak")
    gej = (g.groupby("komoditas_key").agg(gejolak=("gejolak_mingguan", "median"), perubahan=("rata_perubahan_per_minggu", "median"))
           .assign(nama=lambda d: d.index.map(kom.nama_dashboard)))
    top_kota = (g.sort_values("gejolak_mingguan", ascending=False).groupby("komoditas_key").head(1)
                .assign(kota=lambda d: d.kota_key.map(kota), nama=lambda d: d.komoditas_key.map(kom.nama_dashboard)))
    rawit_top = top_kota[top_kota.nama == "Cabai rawit merah"].iloc[0]
    # versi desain menampilkan rata-rata perubahan per minggu per kota (bukan simpangan baku), lihat cek_workbook
    rawit_top_rata = (g[g.komoditas_key.map(kom.nama_dashboard) == "Cabai rawit merah"]
                      .sort_values("rata_perubahan_per_minggu", ascending=False).iloc[0])

    uji = _t("fact_uji_peringatan").set_index("urutan")
    pa = _t("fact_penentu_arah")
    pa = pa[pa.tampil_di_dashboard].assign(kota=lambda d: d.kota_key.map(kota), nama=lambda d: d.komoditas_key.map(kom.nama_dashboard))
    m = _t("fact_margin")
    marg = (m.groupby("komoditas_key").agg(hemat_produsen=("hemat_beli_produsen_pct", "median"),
                                           hemat_pb=("hemat_beli_pedagang_besar_pct", "median"),
                                           margin_pb=("tambahan_di_pedagang_besar_pct", "median"),
                                           margin_pasar=("tambahan_di_pasar_pct", "median"))
            .assign(nama=lambda d: d.index.map(kom.nama_dashboard)))
    bm = kom[kom.nama_dashboard == "Bawang merah"].iloc[0]

    return {
        "minggu_terakhir": akhir,
        "ringkasan": {
            "jml_komoditas_beli_awal": int(tind[tind.beli_lebih_awal].komoditas_key.nunique()),
            "rute_layak_kons": int(dr.layak.sum()), "rute_aman": int((dr.layak & dr.layak_jika_jual_ke_pedagang_besar).sum()),
            "rute_untung_minggu_ini": int((untung_akhir > 0).sum()),
            "naik_tercepat": {"komoditas": tercepat.nama, "perubahan_4m": float(tercepat.perubahan_4m), "harga": float(tercepat.harga)},
            "selisih_nasional_rawit": float(harga[harga.nama == "Cabai rawit merah"].selisih_nasional_pct.iloc[0]),
            "peringatan_aktif_kota": int(al[(al.minggu_key == al.minggu_key.max()) & al.peringatan].kota_key.nunique()),
            "harga_kota_tujuan": harga[["nama", "harga", "perubahan_4m", "harga_nasional", "selisih_nasional_pct"]].round(4).to_dict("records"),
            "tindakan": tind.sort_values(["komoditas_key", "prioritas"])[["komoditas_key", "prioritas", "tindakan", "alasan_singkat"]].to_dict("records"),
            "hal": hal.sort_values("urutan").kalimat.tolist(),
        },
        "rute": {
            "per_pengali": sens, "pengali_acuan": kons,
            "terbaik": {"nama": terbaik.nama_rute, "margin_median": float(terbaik.median_untung_rp_per_kg),
                        "selisih_median": float(terbaik.median_selisih_harga_rp_per_kg), "minggu_untung_dari_10": float(terbaik.minggu_untung_dari_10),
                        "berturut": int(terbaik.minggu_untung_berturut_terakhir), "margin_jual_pb": float(terbaik.median_untung_jual_ke_pedagang_besar)},
            "top10": dr.sort_values("median_untung_rp_per_kg", ascending=False).head(10)[
                ["nama_rute", "median_selisih_harga_rp_per_kg", "median_ongkos_rp_per_kg", "median_untung_rp_per_kg", "minggu_untung_dari_10", "status"]].round(0).to_dict("records"),
        },
        "gejolak": {"per_komoditas": gej.round(4).to_dict("records"), "kota_paling_tidak_stabil": top_kota[["nama", "kota", "gejolak_mingguan"]].round(4).to_dict("records"),
                    "rawit": {"kota": rawit_top.kota, "gejolak": float(rawit_top.gejolak_mingguan)},
                    "rawit_rata": {"kota": kota[rawit_top_rata.kota_key], "gejolak": float(rawit_top_rata.rata_perubahan_per_minggu)}},
        "peringatan": {"benar_dari_10": float(uji.loc[1, "benar_dari_10"]), "tertangkap_dari_10": float(uji.loc[1, "tertangkap_dari_10"]),
                       "lebih_awal_minggu": float(uji.loc[1, "lebih_awal_minggu"]), "benar_tahun_lalu": float(uji.loc[4, "benar_dari_10"]),
                       "tertangkap_tahun_lalu": float(uji.loc[4, "tertangkap_dari_10"]), "benar_selalu": float(uji.loc[5, "benar_dari_10"]),
                       "layak_operasional": bool(uji.loc[1, "benar_dari_10"] > uji.loc[4, "benar_dari_10"]),
                       "kota_naik_lebih_dulu": pa.sort_values(["komoditas_key", "peringkat_penentu_arah"])[["nama", "kota", "jeda_minggu", "hubungan"]].to_dict("records")},
        "pola": {"tabel": kom[["nama_dashboard", "idulfitri_mulai_naik_minggu", "idulfitri_kenaikan_minggu_min1", "idulfitri_tahun_konsisten",
                               "idulfitri_puncak_minggu", "idulfitri_puncak_kenaikan"]].round(4).to_dict("records"),
                 "bawang_merah_mulai_naik": int(bm.idulfitri_mulai_naik_minggu), "bawang_merah_kenaikan_min1": float(bm.idulfitri_kenaikan_minggu_min1)},
        "margin": {"per_komoditas": marg.round(4).to_dict("records"), "hemat_produsen_maks": float(marg.hemat_produsen.max()),
                   "modern_vs_tradisional": kom[["nama_dashboard", "modern_vs_tradisional_pct"]].round(4).to_dict("records")},
    }


# ---------------------------------------------------------------- bagian 2: cek workbook
def pemetaan_acuan(acuan: dict) -> dict[str, tuple]:
    """nama sel -> (nilai acuan, toleransi, keterangan). Nilai None = cukup ada & tidak error."""
    r, ru, gj, pr, po, mg = acuan["ringkasan"], acuan["rute"], acuan["gejolak"], acuan["peringatan"], acuan["pola"], acuan["margin"]
    kons = ru["pengali_acuan"]
    return {
        "kpi_beli_awal": (r["jml_komoditas_beli_awal"], 0, "komoditas perlu dibeli lebih awal"),
        "kpi_rute_untung": (r["rute_layak_kons"], 0, f"rute layak pada ongkos ×{kons:g}"),
        "kpi_kenaikan_tercepat": (r["naik_tercepat"]["perubahan_4m"], 0.005, f"kenaikan 4 minggu {r['naik_tercepat']['komoditas']}"),
        "kpi_peringatan_aktif": (r["peringatan_aktif_kota"], 0, "kota dengan peringatan aktif"),
        "kpi_selisih_nasional_rawit": (r["selisih_nasional_rawit"], 0.005, "cabai rawit kota tujuan vs nasional"),
        "pilih_komoditas": ("Cabai rawit merah", None, "pilihan awal dropdown"),
        "kpi_kota_paling_tidak_stabil": (gj["rawit"]["kota"], None, "kota paling tidak stabil, cabai rawit"),
        "kpi_gejolak_kota_teratas": (gj["rawit"]["gejolak"], 0.0005, "gejolak kota teratas, cabai rawit"),
        "pengali_ongkos": (kons, 0, "pengali ongkos awal"),
        "tambahan_ongkos": (0, 0, "tambahan ongkos awal"),
        "kpi_rute_untung_dinamis": (ru["per_pengali"][kons]["layak"], 0, f"rute untung (dinamis) pada ×{kons:g}"),
        "kpi_rute_aman": (r["rute_aman"], 0, "rute aman"),
        "kpi_rute_margin_positif": (ru["per_pengali"][kons]["median_positif"], 0, f"rute margin positif pada ×{kons:g}"),
        "kpi_margin_rute_terbaik": (ru["terbaik"]["margin_median"], 1, "margin median rute terbaik"),
        "kpi_benar_dari_10": (pr["benar_dari_10"], 0.05, "peringatan benar dari 10"),
        "kpi_tertangkap_dari_10": (pr["tertangkap_dari_10"], 0.05, "lonjakan tertangkap dari 10"),
        "kpi_lebih_awal": (pr["lebih_awal_minggu"], 0.05, "lebih awal (minggu)"),
        "kpi_benar_tahun_lalu": (pr["benar_tahun_lalu"], 0.05, "pola tahun lalu benar dari 10"),
        "kpi_bawang_merah_mulai_naik": (po["bawang_merah_mulai_naik"], 0, "bawang merah mulai naik (minggu)"),
        "kpi_hemat_produsen_maks": (mg["hemat_produsen_maks"], 0.005, "hemat beli di produsen terbesar"),
    }


def baca_nama(path: Path) -> tuple[dict[str, object], dict[str, str]]:
    """Nilai cache sel bernama (workbook harus sudah dihitung ulang)."""
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    nilai, ref = {}, {}
    for name, dn in wb.defined_names.items():
        for sheet, coord in dn.destinations:
            nilai[name] = wb[sheet][coord.replace("$", "")].value
            ref[name] = f"{sheet}!{coord}"
    return nilai, ref


def cek_workbook(path: Path, acuan: dict) -> dict:
    from openpyxl import load_workbook
    nilai, ref = baca_nama(path)
    peta = pemetaan_acuan(acuan)
    if nilai.get("ukuran_gejolak") == "rata_perubahan":   # workbook menyatakan ukuran gejolak yang ditampilkannya
        rr = acuan["gejolak"]["rawit_rata"]
        peta["kpi_kota_paling_tidak_stabil"] = (rr["kota"], None, "kota paling tidak stabil, cabai rawit (rata-rata perubahan)")
        peta["kpi_gejolak_kota_teratas"] = (rr["gejolak"], 0.0005, "rata-rata perubahan per minggu kota teratas, cabai rawit")
    baris = []
    for name, (target, tol, ket) in peta.items():
        v = nilai.get(name, "<tidak ada>")
        if tol is None:
            cocok = (v == target) if isinstance(target, str) else v is not None
        else:
            cocok = isinstance(v, (int, float)) and abs(float(v) - float(target)) <= tol
        baris.append({"nama": name, "sel": ref.get(name, ""), "di_dashboard": v, "acuan": target, "cocok": "✅" if cocok else "❌", "keterangan": ket})
    # error rumus & istilah teknis (nilai cache; sel rumus kosong = belum dihitung ulang)
    wb_v = load_workbook(path, data_only=True)
    wb_f = load_workbook(path)
    errors, istilah, belum_hitung = [], [], 0
    for sheet in SHEET_DASHBOARD:
        wsv, wsf = wb_v[sheet], wb_f[sheet]
        for row in wsf.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    v = wsv[c.coordinate].value
                    if v is None:
                        belum_hitung += 1
                    elif isinstance(v, str) and v.startswith("#") and v != "#N/A":
                        errors.append(f"{sheet}!{c.coordinate}: {v}")
                v = wsv[c.coordinate].value
                if isinstance(v, str) and c.column <= 14 and any(k in v.lower() for k in ISTILAH_TERLARANG):
                    istilah.append(f"{sheet}!{c.coordinate}: {v[:80]}")
    return {"baris": baris, "errors": errors, "istilah": istilah, "belum_dihitung": belum_hitung,
            "cocok": sum(b["cocok"] == "✅" for b in baris), "total": len(baris)}


# ---------------------------------------------------------------- laporan
def _md(d: list[dict]) -> str:
    if not d:
        return "_(kosong)_"
    cols = list(d[0].keys())
    lines = ["| " + " | ".join(cols) + " |", "|" + "---|" * len(cols)]
    lines += ["| " + " | ".join(str(r.get(c, "")) for c in cols) + " |" for r in d]
    return "\n".join(lines)


def write_report(acuan: dict, cek: dict | None, path: Path, xlsx: Path | None) -> None:
    r, ru, gj, pr = acuan["ringkasan"], acuan["rute"], acuan["gejolak"], acuan["peringatan"]
    sens = " · ".join(f"×{p:g}: {v['median_positif']} (layak {v['layak']})" for p, v in ru["per_pengali"].items())
    bag = ""
    if cek is not None:
        bag = f"""
## Cek workbook `{xlsx.as_posix() if xlsx else ''}`
- Nilai sel bernama cocok dengan acuan: **{cek['cocok']} dari {cek['total']}**
- Sel rumus bererror: **{len(cek['errors'])}**{(' — ' + '; '.join(cek['errors'][:20])) if cek['errors'] else ''}
- Sel rumus belum dihitung (jalankan `python -m src.recalc_excel`): {cek['belum_dihitung']}
- Teks dengan istilah teknis terlarang: **{len(cek['istilah'])}**{(' — ' + '; '.join(cek['istilah'][:10])) if cek['istilah'] else ''}

{_md(cek['baris'])}
"""
    text = f"""# Audit dashboard Excel

Dibuat otomatis oleh `python -m src.audit_dashboard` pada {datetime.now():%Y-%m-%d %H:%M}. Nilai acuan dihitung langsung dari
`data/processed/` (minggu terakhir `{acuan['minggu_terakhir']}`), lalu dibandingkan dengan nilai sel bernama di workbook.
{bag}
## Nilai acuan — Ringkasan
- Komoditas perlu dibeli lebih awal: **{r['jml_komoditas_beli_awal']}**
- Rute untung (ongkos ×{ru['pengali_acuan']:g}): **{r['rute_layak_kons']} dari 32**; rute aman: **{r['rute_aman']}**; untung minggu ini: **{r['rute_untung_minggu_ini']}**
- Kenaikan tercepat: **{r['naik_tercepat']['komoditas']} {r['naik_tercepat']['perubahan_4m']:+.0%}** (Rp {r['naik_tercepat']['harga']:,.0f}/kg)
- Cabai rawit kota tujuan vs nasional: **{r['selisih_nasional_rawit']:+.0%}**
- Peringatan aktif: **{r['peringatan_aktif_kota']} kota**

{_md(r['harga_kota_tujuan'])}

- Tindakan (fact_tindakan):

{_md(r['tindakan'])}

- Tiga hal (fact_hal):
{chr(10).join(f'  {i + 1}. {k}' for i, k in enumerate(r['hal']))}

## Nilai acuan — Peluang pindah stok
- Rute dengan median margin > 0 per pengali: {sens}
- Rute terbaik: **{ru['terbaik']['nama']}** — margin median Rp {ru['terbaik']['margin_median']:,.0f}/kg, selisih Rp {ru['terbaik']['selisih_median']:,.0f},
  untung {ru['terbaik']['minggu_untung_dari_10']:.1f} dari 10 minggu, {ru['terbaik']['berturut']} minggu berturut-turut, jika jual ke pedagang besar Rp {ru['terbaik']['margin_jual_pb']:,.0f}

{_md(ru['top10'])}

## Nilai acuan — Gejolak
{_md(gj['per_komoditas'])}

{_md(gj['kota_paling_tidak_stabil'])}

## Nilai acuan — Peringatan dini
- Benar {pr['benar_dari_10']:.1f} dari 10 · tertangkap {pr['tertangkap_dari_10']:.1f} · lebih awal {pr['lebih_awal_minggu']:.1f} minggu ·
  pola tahun lalu benar {pr['benar_tahun_lalu']:.1f} (tertangkap {pr['tertangkap_tahun_lalu']:.1f}) · selalu {pr['benar_selalu']:.1f}
- **Layak operasional: {'YA' if pr['layak_operasional'] else 'TIDAK'}** → keputusan di dashboard harus "{'boleh dipakai sebagai pemicu cek' if pr['layak_operasional'] else 'jangan jadikan pemicu pembelian'}"

{_md(pr['kota_naik_lebih_dulu'])}

## Nilai acuan — Pola musiman
{_md(acuan['pola']['tabel'])}

## Nilai acuan — Rantai pasok
{_md(acuan['margin']['per_komoditas'])}

{_md(acuan['margin']['modern_vs_tradisional'])}

## Pemeriksaan manual (keputusan tidak menyesatkan)
| Halaman | Keputusan | Angka pendukung tampil di halaman yang sama? | Batasan disebut? | Tidak melebihi temuan Q1–Q6? |
|---|---|---|---|---|
| | | | | |
"""
    path.write_text(text, encoding="utf-8")


def main() -> None:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--xlsx", help="workbook dashboard yang sudah dihitung ulang")
    a = p.parse_args()
    cfg, lcfg = load_cfg()
    acuan = nilai_acuan(cfg, lcfg)
    cek = cek_workbook(Path(a.xlsx), acuan) if a.xlsx else None
    out = ROOT / "docs" / "audit_dashboard_report.md"
    write_report(acuan, cek, out, Path(a.xlsx) if a.xlsx else None)
    print(f"Laporan: {out.relative_to(ROOT)}")
    if cek:
        print(f"Cocok: {cek['cocok']}/{cek['total']} · error rumus: {len(cek['errors'])} · belum dihitung: {cek['belum_dihitung']} · istilah terlarang: {len(cek['istilah'])}")
        for b in cek["baris"]:
            if b["cocok"] == "❌":
                print("  ❌", b["nama"], "dashboard =", b["di_dashboard"], "| acuan =", b["acuan"])


if __name__ == "__main__":
    main()
